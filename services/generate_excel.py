
import io
import pandas as pd
from datetime import datetime
from aiogram.types import BufferedInputFile
from sqlalchemy.orm import Session
from database import database
from database.tables import *
from sqlalchemy.orm import Session



def generate_excel_report(chat_id: int) -> BufferedInputFile:

    import io
    import pandas as pd

    from datetime import datetime
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Font,
        PatternFill,
        Border,
        Side,
        Alignment,
    )
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo

    with Session(database.engine) as session:

        # ---------------------------------------------------------
        # 1. Пользователь
        # ---------------------------------------------------------

        user = (
            session.query(Users)
            .filter(Users.chat_id == chat_id)
            .first()
        )

        if user is None:
            raise ValueError(f"Пользователь {chat_id} не найден")

        # ---------------------------------------------------------
        # 2. Операции
        # ---------------------------------------------------------

        records = (
            session.query(
                Actions.date,
                Actions.value,
                Actions.desc,
                Status.type,
                Category.name,
            )
            .join(Status, Actions.status_id == Status.id)
            .join(Category, Actions.category_id == Category.id)
            .filter(Actions.chat_id == chat_id)
            .order_by(Actions.date.desc())
            .all()
        )

        # ---------------------------------------------------------
        # 3. Данные операций
        # ---------------------------------------------------------

        operations = []

        for record in records:

            operation_type = record.type or "Неизвестно"


            value = float(record.value or 0)

            if "списание" in operation_type.lower():
                value = -abs(value)

            operations.append({
                "Дата и время": record.date,
                "Тип": operation_type,
                "Категория": record.name or "Без категории",
                "Сумма": value,
                "Описание": record.desc or "",
            })

        df_operations = pd.DataFrame(
            operations,
            columns=[
                "Дата и время",
                "Тип",
                "Категория",
                "Сумма",
                "Описание",
            ],
        )

        # ---------------------------------------------------------
        # 4. Считаем финансовые показатели
        # ---------------------------------------------------------

        income = sum(
            row["Сумма"]
            for row in operations
            if row["Сумма"] > 0
        )

        expense = abs(sum(
            row["Сумма"]
            for row in operations
            if row["Сумма"] < 0
        ))

        balance_change = income - expense

        # ---------------------------------------------------------
        # 5. MonthlySummary
        # ---------------------------------------------------------

        summary = (
            session.query(MountlySummary)
            .filter(MountlySummary.chat_id == chat_id)
            .order_by(
                MountlySummary.year.desc(),
                MountlySummary.month.desc(),
            )
            .first()
        )

        if summary:
            start_balance = float(summary.start_balance or 0)
            end_balance = float(summary.end_balance or 0)

            if summary.total_income is not None:
                income = float(summary.total_income)

            if summary.total_expense is not None:
                expense = float(summary.total_expense)

        else:
            start_balance = 0
            end_balance = balance_change

        # ---------------------------------------------------------
        # 6. Период
        # ---------------------------------------------------------

        if records:
            dates = [record.date for record in records]

            period_start = min(dates)
            period_end = max(dates)

            period_text = (
                f"{period_start.strftime('%d.%m.%Y')} — "
                f"{period_end.strftime('%d.%m.%Y')}"
            )
        else:
            period_text = "Нет операций"

        # ---------------------------------------------------------
        # 7. Аналитика расходов и доходов
        # ---------------------------------------------------------

        # ---------- Расходы ----------

        df_expenses = df_operations[
            df_operations["Сумма"] < 0
        ].copy()

        if not df_expenses.empty:
            expense_analytics = (
                df_expenses
                .assign(
                    Сумма=df_expenses["Сумма"].abs()
                )
                .groupby(
                    "Категория",
                    as_index=False
                )["Сумма"]
                .sum()
                .sort_values(
                    "Сумма",
                    ascending=False
                )
            )
        else:
            expense_analytics = pd.DataFrame(
                columns=["Категория", "Сумма"]
            )

        # ---------- Доходы ----------

        df_income = df_operations[
            df_operations["Сумма"] > 0
        ].copy()

        if not df_income.empty:
            income_analytics = (
                df_income
                .groupby(
                    "Категория",
                    as_index=False
                )["Сумма"]
                .sum()
                .sort_values(
                    "Сумма",
                    ascending=False
                )
            )
        else:
            income_analytics = pd.DataFrame(
                columns=["Категория", "Сумма"]
            )

        # ---------------------------------------------------------
        # 8. Создаём Excel в памяти
        # ---------------------------------------------------------

        output = io.BytesIO()

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:

            # ==============================================
            # ЛИСТ 1 — СВОДКА
            # ==============================================

            summary_data = pd.DataFrame({
                "Показатель": [
                    "Пользователь",
                    "Username",
                    "Дата регистрации",
                    "Период операций",
                    "",
                    "Начальный баланс",
                    "Доходы",
                    "Расходы",
                    "Изменение баланса",
                    "Конечный баланс",
                ],
                "Значение": [
                    " ".join(
                        x for x in [
                            user.name,
                            user.second_name
                        ]
                        if x
                    ) or "Не указано",

                    f"@{user.username}"
                    if user.username
                    else "Не указан",

                    user.created_at.strftime("%d.%m.%Y")
                    if user.created_at
                    else "—",

                    period_text,

                    "",

                    start_balance,
                    income,
                    expense,
                    income - expense,
                    end_balance,
                ],
            })

            summary_data.to_excel(
                writer,
                index=False,
                sheet_name="Сводка",
                startrow=2,
            )

            # ==============================================
            # ЛИСТ 2 — ОПЕРАЦИИ
            # ==============================================

            df_operations.to_excel(
                writer,
                index=False,
                sheet_name="Операции",
            )

            # ==============================================
            # ЛИСТ 3 — АНАЛИТИКА РАСХОДОВ
            # ==============================================

            expense_analytics.to_excel(
                writer,
                index=False,
                sheet_name="Аналитика расходов",
            )

            # ==============================================
            # ЛИСТ 4 — АНАЛИТИКА ДОХОДОВ
            # ==============================================

            income_analytics.to_excel(
                writer,
                index=False,
                sheet_name="Аналитика доходов",
            )

        # ---------------------------------------------------------
        # 9. Открываем созданный Excel для оформления
        # ---------------------------------------------------------

        output.seek(0)

        workbook = load_workbook(output)

        ws_summary = workbook["Сводка"]
        ws_operations = workbook["Операции"]
        ws_expenses = workbook["Аналитика расходов"]
        ws_income = workbook["Аналитика доходов"]

        # ---------------------------------------------------------
        # Цвета
        # ---------------------------------------------------------

        BLUE = "4472C4"
        DARK_BLUE = "1F4E78"
        LIGHT_BLUE = "D9EAF7"

        GREEN = "70AD47"
        LIGHT_GREEN = "E2F0D9"

        RED = "C00000"
        LIGHT_RED = "FCE4D6"

        GRAY = "D9E1F2"
        WHITE = "FFFFFF"

        # ---------------------------------------------------------
        # Общие стили
        # ---------------------------------------------------------

        thin_border = Border(
            left=Side(
                style="thin",
                color="D9E1F2"
            ),
            right=Side(
                style="thin",
                color="D9E1F2"
            ),
            top=Side(
                style="thin",
                color="D9E1F2"
            ),
            bottom=Side(
                style="thin",
                color="D9E1F2"
            ),
        )

        # =========================================================
        # СВОДКА
        # =========================================================

        ws_summary.merge_cells("A1:B1")

        title = ws_summary["A1"]
        title.value = "ФИНАНСОВЫЙ ОТЧЁТ"

        title.font = Font(
            size=20,
            bold=True,
            color=WHITE,
        )

        title.fill = PatternFill(
            "solid",
            fgColor=DARK_BLUE,
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        ws_summary.row_dimensions[1].height = 35


        for cell in ws_summary[3]:
            cell.font = Font(
                bold=True,
                color=WHITE,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=BLUE,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in range(
            4,
            ws_summary.max_row + 1
        ):
            ws_summary.cell(
                row,
                1
            ).border = thin_border

            ws_summary.cell(
                row,
                2
            ).border = thin_border

            ws_summary.cell(
                row,
                1
            ).font = Font(
                bold=True
            )

        for row in range(9, 14):
            ws_summary.cell(
                row,
                2
            ).number_format = '#,##0.00 "₽"'

        ws_summary["B10"].fill = PatternFill(
            "solid",
            fgColor=LIGHT_GREEN,
        )

        ws_summary["B10"].font = Font(
            bold=True,
            color=GREEN,
        )

        ws_summary["B11"].fill = PatternFill(
            "solid",
            fgColor=LIGHT_RED,
        )

        ws_summary["B11"].font = Font(
            bold=True,
            color=RED,
        )

        ws_summary["B13"].font = Font(
            size=14,
            bold=True,
        )

        ws_summary.column_dimensions["A"].width = 28
        ws_summary.column_dimensions["B"].width = 35

        # =========================================================
        # ОПЕРАЦИИ
        # =========================================================

        ws_operations.freeze_panes = "A2"

        ws_operations.auto_filter.ref = (
            ws_operations.dimensions
        )

        for cell in ws_operations[1]:
            cell.font = Font(
                bold=True,
                color=WHITE,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=BLUE,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in range(
            2,
            ws_operations.max_row + 1
        ):
            date_cell = ws_operations.cell(
                row,
                1
            )

            amount_cell = ws_operations.cell(
                row,
                4
            )

            date_cell.number_format = (
                "dd.mm.yyyy hh:mm"
            )

            amount_cell.number_format = (
                '#,##0.00 "₽"'
            )

            for cell in ws_operations[row]:
                cell.border = thin_border

            if (
                amount_cell.value is not None
                and amount_cell.value > 0
            ):
                amount_cell.font = Font(
                    bold=True,
                    color=GREEN,
                )

                amount_cell.fill = PatternFill(
                    "solid",
                    fgColor=LIGHT_GREEN,
                )

            elif (
                amount_cell.value is not None
                and amount_cell.value < 0
            ):
                amount_cell.font = Font(
                    bold=True,
                    color=RED,
                )

                amount_cell.fill = PatternFill(
                    "solid",
                    fgColor=LIGHT_RED,
                )

        if ws_operations.max_row >= 2:

            table_ref = (
                f"A1:E{ws_operations.max_row}"
            )

            table = Table(
                displayName="OperationsTable",
                ref=table_ref,
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style

            ws_operations.add_table(table)

        # Ширина колонок
        widths = {
            "A": 20,
            "B": 18,
            "C": 25,
            "D": 18,
            "E": 50,
        }

        for column, width in widths.items():
            ws_operations.column_dimensions[
                column
            ].width = width

        # =========================================================
        # АНАЛИТИКА РАСХОДОВ
        # =========================================================

        ws_expenses.freeze_panes = "A2"

        for cell in ws_expenses[1]:
            cell.font = Font(
                bold=True,
                color=WHITE,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=RED,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in range(
            2,
            ws_expenses.max_row + 1
        ):
            ws_expenses.cell(
                row,
                2
            ).number_format = '#,##0.00 "₽"'

            for cell in ws_expenses[row]:
                cell.border = thin_border

        ws_expenses.column_dimensions["A"].width = 30
        ws_expenses.column_dimensions["B"].width = 20

        # Диаграмма расходов
        if ws_expenses.max_row >= 2:

            expense_chart = BarChart()

            expense_chart.type = "bar"
            expense_chart.style = 10

            expense_chart.title = (
                "Расходы по категориям"
            )

            expense_chart.y_axis.title = (
                "Категория"
            )

            expense_chart.x_axis.title = (
                "Сумма, ₽"
            )

            data = Reference(
                ws_expenses,
                min_col=2,
                min_row=1,
                max_row=ws_expenses.max_row,
            )

            categories = Reference(
                ws_expenses,
                min_col=1,
                min_row=2,
                max_row=ws_expenses.max_row,
            )

            expense_chart.add_data(
                data,
                titles_from_data=True,
            )

            expense_chart.set_categories(
                categories
            )

            expense_chart.height = 8
            expense_chart.width = 16

            ws_expenses.add_chart(
                expense_chart,
                "D2",
            )

        # =========================================================
        # АНАЛИТИКА ДОХОДОВ
        # =========================================================

        ws_income.freeze_panes = "A2"

        for cell in ws_income[1]:
            cell.font = Font(
                bold=True,
                color=WHITE,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=GREEN,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row in range(
            2,
            ws_income.max_row + 1
        ):
            ws_income.cell(
                row,
                2
            ).number_format = '#,##0.00 "₽"'

            for cell in ws_income[row]:
                cell.border = thin_border

        ws_income.column_dimensions["A"].width = 30
        ws_income.column_dimensions["B"].width = 20

        # ДИАГРАММА ДОХОДОВ
        # =========================================================

        if ws_income.max_row >= 2:

            income_chart = BarChart()

            income_chart.type = "bar"

            income_chart.grouping = "clustered"
            income_chart.overlap = 0

            income_chart.title = "Доходы по категориям"

            income_chart.y_axis.title = "Категория"

            income_chart.x_axis.numFmt = '#,##0 "₽"'

            data = Reference(
                ws_income,
                min_col=2,
                min_row=1,
                max_row=ws_income.max_row,
            )

            categories = Reference(
                ws_income,
                min_col=1,
                min_row=2,
                max_row=ws_income.max_row,
            )

            income_chart.add_data(
                data,
                titles_from_data=True,
            )

            income_chart.set_categories(
                categories
            )

            # Размер
            income_chart.height = 8
            income_chart.width = 15

            # Расстояние между полосами
            income_chart.gapWidth = 50

            # Добавляем диаграмму
            ws_income.add_chart(
                income_chart,
                "D2",
            )

        # ---------------------------------------------------------
        # 10. Сохраняем workbook обратно в RAM
        # ---------------------------------------------------------

        final_output = io.BytesIO()

        workbook.save(final_output)

        final_output.seek(0)

        # ---------------------------------------------------------
        # 11. Возвращаем aiogram-файл
        # ---------------------------------------------------------

        filename = (
            f"financial_report_"
            f"{datetime.now().strftime('%Y_%m_%d')}.xlsx"
        )

        return BufferedInputFile(
            final_output.read(),
            filename=filename,
        )